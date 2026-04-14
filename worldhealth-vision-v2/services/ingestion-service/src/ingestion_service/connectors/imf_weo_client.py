import io
from datetime import UTC, datetime

import httpx
from openpyxl import load_workbook
from tenacity import retry, retry_if_exception_type, stop_after_attempt, wait_exponential


class ImfWeoClient:
    def __init__(self, workbook_url: str, timeout_seconds: int) -> None:
        self._workbook_url = workbook_url
        self._timeout_seconds = timeout_seconds

    @retry(
        reraise=True,
        stop=stop_after_attempt(3),
        wait=wait_exponential(multiplier=1, min=1, max=8),
        retry=retry_if_exception_type(httpx.HTTPError),
    )
    def fetch_countries_sheet_rows(self) -> list[dict[str, object | None]]:
        with httpx.Client(timeout=self._timeout_seconds, follow_redirects=True) as client:
            response = client.get(self._workbook_url)
            response.raise_for_status()
            workbook_bytes = response.content

        workbook = load_workbook(filename=io.BytesIO(workbook_bytes), read_only=True, data_only=True)
        try:
            sheet = workbook["Countries"]
            header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
            headers = [str(value).strip() if value is not None else "" for value in header_row]

            rows: list[dict[str, object | None]] = []
            for values in sheet.iter_rows(min_row=2, values_only=True):
                row = {
                    header: value
                    for header, value in zip(headers, values, strict=False)
                    if header
                }
                if row:
                    rows.append(row)
            return rows
        finally:
            workbook.close()

    def describe_source(self) -> dict[str, object]:
        return {
            "workbookUrl": self._workbook_url,
            "downloadedAtUtc": datetime.now(UTC).isoformat(),
        }
