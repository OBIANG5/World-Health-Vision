import csv
import io
import re
from datetime import UTC, datetime
from pathlib import Path

import httpx
from openpyxl import load_workbook
from tenacity import retry, retry_if_exception_type, stop_after_attempt, wait_exponential


class UnodcTabularClient:
    def __init__(
        self,
        *,
        download_url: str | None,
        local_export_path: Path | None,
        timeout_seconds: int,
    ) -> None:
        self._download_url = download_url
        self._local_export_path = local_export_path
        self._timeout_seconds = timeout_seconds

    @retry(
        reraise=True,
        stop=stop_after_attempt(3),
        wait=wait_exponential(multiplier=1, min=1, max=8),
        retry=retry_if_exception_type(httpx.HTTPError),
    )
    def fetch_dataset(
        self,
        *,
        required_column_alias_groups: list[list[str]],
    ) -> tuple[list[dict[str, str | None]], dict[str, object]]:
        payload_bytes, descriptor = self._load_payload()
        rows = self._parse_payload(
            payload_bytes=payload_bytes,
            descriptor=descriptor,
            required_column_alias_groups=required_column_alias_groups,
        )
        descriptor["rowCount"] = len(rows)
        return rows, descriptor

    def _load_payload(self) -> tuple[bytes, dict[str, object]]:
        if self._local_export_path is not None:
            if not self._local_export_path.exists():
                raise FileNotFoundError(f"UNODC export file not found: {self._local_export_path}")
            payload = self._local_export_path.read_bytes()
            return payload, {
                "mode": "LOCAL_EXPORT",
                "localExportPath": str(self._local_export_path),
                "resolvedFormat": self._infer_format_from_name(self._local_export_path.name),
                "downloadedAtUtc": datetime.now(UTC).isoformat(),
            }

        if not self._download_url:
            raise ValueError(
                "UNODC homicide export is not configured. Set WHV_INGEST_UNODC_HOMICIDE_EXPORT_PATH to an official export file "
                "or provide WHV_INGEST_UNODC_HOMICIDE_DOWNLOAD_URL once a stable direct download URL is confirmed."
            )

        headers = {
            "User-Agent": (
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                "(KHTML, like Gecko) Chrome/135.0.0.0 Safari/537.36"
            ),
            "Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,text/csv,*/*;q=0.8",
            "Referer": "https://dataunodc.un.org/content/homicide-country-data",
        }

        with httpx.Client(timeout=self._timeout_seconds, follow_redirects=True, headers=headers) as client:
            response = client.get(self._download_url)
            response.raise_for_status()
            payload = response.content
            content_type = response.headers.get("content-type")

        if self._looks_like_html(payload, content_type):
            raise ValueError(
                "UNODC returned HTML instead of a downloadable data file. Use WHV_INGEST_UNODC_HOMICIDE_EXPORT_PATH with an official "
                "export file until a stable direct-download URL is confirmed."
            )

        return payload, {
            "mode": "REMOTE_DOWNLOAD",
            "downloadUrl": self._download_url,
            "contentType": content_type,
            "resolvedFormat": self._infer_format_from_name(self._download_url),
            "downloadedAtUtc": datetime.now(UTC).isoformat(),
        }

    def _parse_payload(
        self,
        *,
        payload_bytes: bytes,
        descriptor: dict[str, object],
        required_column_alias_groups: list[list[str]],
    ) -> list[dict[str, str | None]]:
        resolved_format = descriptor.get("resolvedFormat")
        if resolved_format == "csv":
            return self._parse_csv_rows(payload_bytes)
        if resolved_format == "xlsx":
            return self._parse_workbook_rows(payload_bytes, required_column_alias_groups)
        raise ValueError(f"Unsupported UNODC export format: {resolved_format}")

    def _parse_csv_rows(self, payload_bytes: bytes) -> list[dict[str, str | None]]:
        text = payload_bytes.decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text))
        rows: list[dict[str, str | None]] = []
        for row in reader:
            normalized = {
                (key.strip() if key else ""): self._normalize_cell_to_text(value)
                for key, value in row.items()
                if key
            }
            if any(value is not None for value in normalized.values()):
                rows.append(normalized)
        return rows

    def _parse_workbook_rows(
        self,
        payload_bytes: bytes,
        required_column_alias_groups: list[list[str]],
    ) -> list[dict[str, str | None]]:
        workbook = load_workbook(filename=io.BytesIO(payload_bytes), read_only=True, data_only=True)
        try:
            normalized_groups = [
                {self._normalize_alias(alias) for alias in aliases}
                for aliases in required_column_alias_groups
            ]
            for sheet in workbook.worksheets:
                header_row_index, headers = self._find_header_row(sheet, normalized_groups)
                if header_row_index is None:
                    continue

                rows: list[dict[str, str | None]] = []
                for values in sheet.iter_rows(min_row=header_row_index + 1, values_only=True):
                    row = {
                        header: self._normalize_cell_to_text(value)
                        for header, value in zip(headers, values, strict=False)
                        if header
                    }
                    if any(value is not None for value in row.values()):
                        rows.append(row)
                if rows:
                    return rows

            raise ValueError("Unable to locate a structured data sheet in the UNODC export.")
        finally:
            workbook.close()

    def _find_header_row(
        self,
        sheet,
        normalized_groups: list[set[str]],
    ) -> tuple[int | None, list[str]]:
        for row_index, values in enumerate(sheet.iter_rows(min_row=1, max_row=15, values_only=True), start=1):
            headers = [self._normalize_cell_to_text(value) or "" for value in values]
            normalized_headers = {self._normalize_alias(value) for value in headers if value}
            if normalized_headers and all(group & normalized_headers for group in normalized_groups):
                return row_index, headers
        return None, []

    def _looks_like_html(self, payload_bytes: bytes, content_type: str | None) -> bool:
        if content_type and "html" in content_type.lower():
            return True
        prefix = payload_bytes[:128].lstrip().lower()
        return prefix.startswith(b"<!doctype html") or prefix.startswith(b"<html")

    def _infer_format_from_name(self, value: str) -> str:
        lowered = value.lower()
        if lowered.endswith(".csv"):
            return "csv"
        if lowered.endswith(".xlsx") or lowered.endswith(".xlsm"):
            return "xlsx"
        return "xlsx"

    def _normalize_alias(self, value: str) -> str:
        return re.sub(r"[^a-z0-9]+", " ", value.lower()).strip()

    def _normalize_cell_to_text(self, value: object | None) -> str | None:
        if value is None:
            return None
        cleaned = str(value).strip()
        return cleaned or None
